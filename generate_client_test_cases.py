import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def style_header(ws, row, cols, fill_color="1B2A3D"):
    header_font = Font(name='Calibri', bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'))
    for col in range(1, cols+1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

def style_section_header(ws, row, cols, title, fill_color="C9A84C"):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row=row, column=1)
    cell.value = title
    cell.font = Font(name='Calibri', bold=True, color="1B2A3D", size=13)
    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")

def add_data_row(ws, row, data):
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'))
    data_align = Alignment(vertical="top", wrap_text=True)
    for col, value in enumerate(data, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = value
        cell.font = Font(name='Calibri', size=10)
        cell.alignment = data_align
        cell.border = thin_border

HEADERS = ["TC ID", "Test Scenario", "Test Steps", "Test Data / Inputs", "Expected Result", "Priority", "Status", "Remarks"]
NUM_COLS = len(HEADERS)

wb = openpyxl.Workbook()

# --- Sheet 1: Client Management ---
ws = wb.active
ws.title = "Client Management"
ws.sheet_properties.tabColor = "1B2A3D"

# Column widths
col_widths = [10, 35, 50, 40, 40, 10, 10, 20]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

r = 1
style_section_header(ws, r, NUM_COLS, "CLIENT MANAGEMENT — TEST CASES")
r = 2
style_section_header(ws, r, NUM_COLS, "Section A: Add Client Form & Logic", "E8E4DE")
r = 3
for col, h in enumerate(HEADERS, 1):
    ws.cell(row=r, column=col, value=h)
style_header(ws, r, NUM_COLS)

client_tests_a = [
    ["CL-001", "Open Add Client form", "1. Login as Admin/User\n2. Navigate to Client Management > Add Client", "N/A", "Form loads successfully; Sales Partner dropdown populated with active partners.", "High", "", ""],
    ["CL-002", "Add Client with valid data", "1. Fill Name, City, Mobile, Email, Reference, Sales Partner\n2. Click Save", "Name: John Doe, Mobile: 9090762424, email: john@test.com", "Client record saved; Success message 'Client record is updated successfully.' displayed; Redirects to 'view_clients_list'.", "High", "", ""],
    ["CL-003", "Add Client with existing mobile number", "1. Enter a mobile number that is already in DB\n2. Fill other details\n3. Click Save", "Mobile: 9090762424 (Existing)", "Validation error: 'Mobile already exists' displayed for the mobile field.", "High", "", ""],
    ["CL-004", "Add Client with new city", "1. Enter a city name that doesn't exist in the dropdown/search\n2. Fill other details\n3. Click Save", "CityName: Mars", "System creates a new city record in 'cities' table automatically; Client saved successfully.", "Medium", "", ""],
    ["CL-005", "Add Client with null mandatory fields", "1. Leave Name and Mobile empty\n2. Click Save", "Name: empty, Mobile: empty", "Validation errors displayed for mandatory fields (depending on JS/Spring form validation).", "High", "", ""],
]

r = 4
for test in client_tests_a:
    add_data_row(ws, r, test)
    r += 1

# Section B: Client Listing & Filtering
r += 1
style_section_header(ws, r, NUM_COLS, "Section B: Client Listing & Search Logic", "E8E4DE")
r += 1
for col, h in enumerate(HEADERS, 1):
    ws.cell(row=r, column=col, value=h)
style_header(ws, r, NUM_COLS)
r += 1

client_tests_b = [
    ["CL-006", "Load Client Listing page", "1. Navigate to Client Management > View Clients", "N/A", "Client list loads with pagination; default size is 20 rows (VistaluxConstants.DEFAULT_PAGE_SIZE).", "High", "", ""],
    ["CL-007", "Search Client by Name", "1. Enter partial/full name in search\n2. Click Search", "Name: 'John'", "List filters to show only clients whose name contains 'John'.", "High", "", ""],
    ["CL-008", "Search Client by City", "1. Select a city from filter\n2. Click Search", "City: 'Corbett'", "List filters to show clients only from Corbett.", "High", "", ""],
    ["CL-009", "Filter Client by Active/Inactive status", "1. Select Active = 'Inactive'\n2. Click Search", "Active: False", "List shows only deactivated clients.", "Medium", "", ""],
    ["CL-010", "Pagination check", "1. Navigate through Page 1, Page 2", "Total records > 20", "Next page loads correct set of clients; totalPages count is accurate.", "Medium", "", ""],
    ["CL-011", "Search Client by Sales Partner", "1. Select a Sales Partner from filter dropdown\n2. Click Search", "Sales Partner: 'Digital Marketing'", "List filters to show clients assigned to selected partner.", "High", "", ""],
]

for test in client_tests_b:
    add_data_row(ws, r, test)
    r += 1

# Section C: Edit Client Logic
r += 1
style_section_header(ws, r, NUM_COLS, "Section C: Edit Client Logic", "E8E4DE")
r += 1
for col, h in enumerate(HEADERS, 1):
    ws.cell(row=r, column=col, value=h)
style_header(ws, r, NUM_COLS)
r += 1

client_tests_c = [
    ["CL-012", "Load Edit Client form", "1. On listing page, click 'Edit' for a client", "Valid clientId", "Edit form loads with existing data correctly populated (Name, Mobile, Email, etc.).", "High", "", ""],
    ["CL-013", "Update Client with valid changes", "1. Modify Email and Remarks\n2. Click Update", "Email: updated@test.com", "Record updated in DB; Success message displayed; Redirects to 'view_clients_list'.", "High", "", ""],
    ["CL-014", "Update Client with another's mobile number", "1. Change mobile number to one used by ANOTHER client\n2. Click Update", "Mobile: 9191762424 (Used by Client X)", "Validation error: 'Mobile already exists' displayed.", "High", "", ""],
    ["CL-015", "Update Client mobile to same (own) number", "1. Leave mobile as is or re-type same current mobile\n2. Change Name\n3. Click Update", "Mobile: (current number)", "Should process successfully (skip duplicate check for own record).", "High", "", ""],
]

for test in client_tests_c:
    add_data_row(ws, r, test)
    r += 1

# Section D: View & Export
r += 1
style_section_header(ws, r, NUM_COLS, "Section D: View Details & Export Functionality", "E8E4DE")
r += 1
for col, h in enumerate(HEADERS, 1):
    ws.cell(row=r, column=col, value=h)
style_header(ws, r, NUM_COLS)
r += 1

client_tests_d = [
    ["CL-016", "View Client Details page", "1. On listing page, click 'View' icon/button", "Valid clientId", "Detailed read-only view loads showing Name, Mobile, Email, Reference, Sales Partner Name, and City.", "High", "", ""],
    ["CL-017", "Export filtered list to Excel", "1. Apply filters (e.g. City=Delhi)\n2. Click 'Export Excel'", "Filter: Delhi", "System downloads 'clients.xlsx' containing only filtered records.", "High", "", ""],
    ["CL-018", "Export filtered list to PDF", "1. Apply filters (e.g. B2B=True)\n2. Click 'Export PDF'", "Filter: B2B", "System downloads 'clients.pdf' with correct formatting and data.", "High", "", ""],
    ["CL-019", "AJAX Client Search (Autocomplete)", "1. During quotation creation, start typing client name", "Search: 'John'", "System returns matching client list with sales partner names (Used by getClientList endpoint).", "High", "", ""],
]

for test in client_tests_d:
    add_data_row(ws, r, test)
    r += 1

filename = "/Users/sarthakjain/vistaluxhms_repository/base_vistaluxhms_repository/Client_Management_TestCases.xlsx"
wb.save(filename)
print(f"✅ Saved: {filename}")
