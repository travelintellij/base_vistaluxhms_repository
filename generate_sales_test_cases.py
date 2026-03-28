import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def style_header(ws, row, cols, fill_color="1B2A3D"):
    header_font = Font(name='Calibri', bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'))
    for col in range(1, cols+1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

def style_section_header(ws, row, cols, title, fill_color="C9A84C"):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row=row, column=1)
    cell.value = title
    cell.font = Font(name='Calibri', bold=True, color="1B2A3D", size=13)
    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")

def add_data_row(ws, row, data):
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'))
    data_align = Alignment(vertical="top", wrap_text=True)
    for col, value in enumerate(data, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = value
        cell.font = Font(name='Calibri', size=10)
        cell.alignment = data_align
        cell.border = thin_border

HEADERS = ["TC ID", "Test Scenario", "Test Steps", "Test Data / Inputs", "Expected Result", "Priority", "Status", "Remarks"]
NUM_COLS = len(HEADERS)

wb = openpyxl.Workbook()

# --- Sheet 1: Sales Partner Management ---
ws_sp = wb.active
ws_sp.title = "Sales Partner Management"
ws_sp.sheet_properties.tabColor = "1B2A3D"

col_widths = [10, 35, 50, 40, 40, 10, 10, 20]
for i, w in enumerate(col_widths, 1):
    ws_sp.column_dimensions[get_column_letter(i)].width = w

r = 1
style_section_header(ws_sp, r, NUM_COLS, "SALES PARTNER MANAGEMENT (incl. Client Sync)")
r = 2
style_section_header(ws_sp, r, NUM_COLS, "Section A: Creation & Sync Logic", "E8E4DE")
r = 3
for col, h in enumerate(HEADERS, 1):
    ws_sp.cell(row=r, column=col, value=h)
style_header(ws_sp, r, NUM_COLS)

sp_tests_a = [
    ["SP-001", "Create Sales Partner: Verify automatic Client creation", "1. Add a new Sales Partner\n2. Click Save\n3. Check Client Management list", "Name: 'Global Travels', Mobile: 9988776655", "Sales Partner saved; A new Client 'Global Travels' is automatically created with 'B2B=True' and Mobile mapping.", "High", "", "Logic in createEditSalesPartner"],
    ["SP-002", "Duplicate Mobile Check (Sales Partner)", "1. Attempt to add a Sales Partner with a mobile already in SP table\n2. Click Save", "Mobile: 9090762424 (Existing SP)", "Validation error: 'Mobile already exists in Sales Partner' displayed.", "High", "", "Break logic 1"],
    ["SP-003", "Create SP with new City", "1. Enter a non-existent city name\n2. Click Save", "City: 'Wonderland'", "City 'Wonderland' added to DB; SP saved with this city's ID.", "Medium", "", ""],
    ["SP-004", "Edit SP: Verify Client sync", "1. Edit an existing SP (Change Name/Email)\n2. Click Save\n3. Check Client list", "Name updated from 'Global' to 'Universal'", "Corresponding Client record 'Universal' updated automatically.", "High", "", "Logic in edit_edit_sales_partner"],
    ["SP-005", "Edit SP: Same mobile allowed for self", "1. Edit SP\n2. Keep same mobile but change description\n3. Click Save", "Mobile: (unchanged)", "Success; no duplicate error for own record.", "High", "", ""],
    ["SP-006", "Break Case: Create SP without selecting Rate Type", "1. Leave Rate Type dropdown empty/default\n2. Click Save", "RateType: null", "System should show validation error or not permit save (if mandatory).", "High", "", "Check mandatory constraints"],
]

r = 4
for test in sp_tests_a:
    add_data_row(ws_sp, r, test)
    r += 1

# Section B: Rate Sharing Logic
r += 1
style_section_header(ws_sp, r, NUM_COLS, "Section B: Rate Card Sharing (Email Logic)", "E8E4DE")
r += 1
for col, h in enumerate(HEADERS, 1):
    ws_sp.cell(row=r, column=col, value=h)
style_header(ws_sp, r, NUM_COLS)
r += 1

sp_tests_b = [
    ["SP-007", "Share Rate Card: Valid Single Email", "1. Click 'Share Rates'\n2. Choose sessions\n3. Input valid email\n4. Send", "Email: partner@test.com", "Email sent successfully with rate card template.", "High", "", ""],
    ["SP-008", "Share Rate Card: Multiple Emails (Comma)", "1. Input multiple emails separated by comma\n2. Send", "Email: a@b.com, c@d.com", "All recipients receive the email.", "Medium", "", ""],
    ["SP-009", "Break Case: Invalid Email Format", "1. Input 'not-an-email'\n2. Send", "Email: hello-world", "System catches invalid address exception and shows error message.", "High", "", "Break logic 2"],
    ["SP-010", "Share Rate Card: No sessions selected", "1. Navigate to share form\n2. Click 'Review' without checking any sessions", "Sessions: empty", "Error: 'No sessions were selected!' displayed.", "Medium", "", ""],
    ["SP-011", "Verify Person-wise rates in Email Data", "1. Review rate card before sending", "Standard occupancy rooms", "Rates for Person 1, 2, 3... are correctly mapped from session details DTO.", "High", "", ""],
]

for test in sp_tests_b:
    add_data_row(ws_sp, r, test)
    r += 1

# --- Sheet 2: Rate Type & Room Management ---
ws_rm = wb.create_sheet("Rate & Room Management")
ws_rm.sheet_properties.tabColor = "C9A84C"
for i, w in enumerate(col_widths, 1):
    ws_rm.column_dimensions[get_column_letter(i)].width = w

r = 1
style_section_header(ws_rm, r, NUM_COLS, "RATE TYPE & ROOM CATEGORY MANAGEMENT")
r = 2
style_section_header(ws_rm, r, NUM_COLS, "Section A: Rate Type Logic", "E8E4DE")
r = 3
for col, h in enumerate(HEADERS, 1):
    ws_rm.cell(row=r, column=col, value=h)
style_header(ws_rm, r, NUM_COLS)
r = 4

rate_tests = [
    ["RR-001", "Add New Rate Type", "1. Add a unique Rate Type name\n2. Click Save", "Name: 'Corporate Platinum'", "New rate type appears in lists and dropdowns for Sales Partners.", "High", "", ""],
    ["RR-002", "Edit Rate Type: Deactivate", "1. Edit active rate type\n2. Set Active=False\n3. Save", "Active: False", "Deactivated rate type no longer appears in 'Add Sales Partner' dropdown.", "High", "", "Check findAllActiveRateTypes"],
    ["RR-003", "Break Case: Empty Rate Type Name", "1. Attempt to save rate type with empty name", "Name: ''", "Validation error or blocked by DB (Not Null).", "Medium", "", ""],
]

for test in rate_tests:
    add_data_row(ws_rm, r, test)
    r += 1

r += 1
style_section_header(ws_rm, r, NUM_COLS, "Section B: Room Category Logic", "E8E4DE")
r += 1
for col, h in enumerate(HEADERS, 1):
    ws_rm.cell(row=r, column=col, value=h)
style_header(ws_rm, r, NUM_COLS)
r += 1

room_tests = [
    ["RR-004", "Add Room: Extra Bed validation", "1. Add room with Extra Bed occupancy > 0\n2. Click Save", "Name: 'Luxury Suite', ExtraBed: 1", "Room saved; Extra bed logic works in quotations.", "High", "", ""],
    ["RR-005", "Break Case: Max occupancy < Standard occupancy", "1. Set Standard=3, Max=2\n2. Click Save", "Std: 3, Max: 2", "Should be logically blocked (Max must be >= Standard).", "High", "", "Break logic 3"],
    ["RR-006", "Room List Pagination", "1. View rooms listing", "N/A", "Verify all categories appear; typically hotel room count is low so usually on 1 page.", "Low", "", ""],
    ["RR-007", "Break Case: Negative Percentage Values", "1. Set extraBedPercentage = -10\n2. Save", "ExtraBed%: -10", "System should block negative values to prevent pricing errors.", "High", "", "Break logic 4"],
]

for test in room_tests:
    add_data_row(ws_rm, r, test)
    r += 1

filename = "/Users/sarthakjain/vistaluxhms_repository/base_vistaluxhms_repository/Sales_Management_TestCases.xlsx"
wb.save(filename)
print(f"✅ Saved: {filename}")
