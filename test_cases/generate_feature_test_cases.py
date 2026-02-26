"""
Generate SHORTLISTED test case Excel & CSV files for:
  1. Add Client (essential mandatory tests only)
  2. Rate Type Management (Add / Edit / List)

Each feature gets its own Excel + CSV file.
"""
import csv
import os
import sys

# Use the venv's openpyxl
for subdir in os.listdir("/tmp/testcase_venv/lib"):
    site_packages = f"/tmp/testcase_venv/lib/{subdir}/site-packages"
    if os.path.isdir(site_packages):
        sys.path.insert(0, site_packages)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ───────────────── Style definitions ──────────────────
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
CATEGORY_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
CATEGORY_FONT = Font(name="Calibri", bold=True, size=11, color="1F4E79")
DATA_FONT = Font(name="Calibri", size=10)
PRIORITY_FILLS = {
    "Critical": PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid"),
    "High":     PatternFill(start_color="FF9900", end_color="FF9900", fill_type="solid"),
    "Medium":   PatternFill(start_color="FFCC00", end_color="FFCC00", fill_type="solid"),
    "Low":      PatternFill(start_color="66CC66", end_color="66CC66", fill_type="solid"),
}
PRIORITY_FONTS = {
    "Critical": Font(name="Calibri", bold=True, size=10, color="FFFFFF"),
    "High":     Font(name="Calibri", bold=True, size=10, color="000000"),
    "Medium":   Font(name="Calibri", bold=True, size=10, color="000000"),
    "Low":      Font(name="Calibri", bold=True, size=10, color="000000"),
}
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)

HEADERS = [
    "TC ID", "Category", "Priority", "Test Case Title",
    "Pre-Conditions", "Test Steps", "Test Data",
    "Expected Result", "Status", "Remarks"
]
COL_WIDTHS = [12, 22, 10, 40, 28, 45, 40, 50, 12, 25]


def write_excel(test_cases, filepath, sheet_title):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    # Header row
    for c, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for c, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    row = 2
    last_category = None
    for tc in test_cases:
        # Category separator
        if tc["category"] != last_category:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(HEADERS))
            cell = ws.cell(row=row, column=1, value=tc["category"])
            cell.fill = CATEGORY_FILL
            cell.font = CATEGORY_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")
            for cc in range(1, len(HEADERS) + 1):
                ws.cell(row=row, column=cc).border = THIN_BORDER
            row += 1
            last_category = tc["category"]

        values = [
            tc["id"], tc["category"], tc["priority"], tc["title"],
            tc["preconditions"], tc["steps"], tc["data"],
            tc["expected"], "", ""
        ]
        for c, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=val)
            cell.font = DATA_FONT
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = THIN_BORDER
            if c == 3:  # Priority column
                cell.fill = PRIORITY_FILLS.get(val, PatternFill())
                cell.font = PRIORITY_FONTS.get(val, DATA_FONT)
                cell.alignment = Alignment(horizontal="center", vertical="top")
        row += 1

    wb.save(filepath)
    return len(test_cases)


def write_csv(test_cases, filepath):
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(HEADERS)
        for tc in test_cases:
            writer.writerow([
                tc["id"], tc["category"], tc["priority"], tc["title"],
                tc["preconditions"], tc["steps"], tc["data"],
                tc["expected"], "", ""
            ])


# ═══════════════════════════════════════════════════════════════════════════
#  ADD CLIENT — SHORTLISTED ESSENTIAL TEST CASES (Manual Testing)
# ═══════════════════════════════════════════════════════════════════════════
ADD_CLIENT_TESTS = [
    # ── Page Load ──
    {
        "id": "TC_AC_001", "category": "Page Load & Navigation", "priority": "High",
        "title": "Verify Add Client page loads with all fields visible",
        "preconditions": "User is logged in as ADMIN",
        "steps": "1. Navigate to Sales Management > Client Management\n2. Click 'Add Client'",
        "data": "URL: /ashokacrm/view_add_client_form",
        "expected": "Page loads with fields: Client Id (Auto Generated in blue), Client Source dropdown, Client Name, Client Type (B2B/B2C), Email Id, Mobile, City (autocomplete), Active dropdown, Reference, Remarks. Buttons: 'Add Client' and 'View Client List'."
    },
    # ── Happy Path ──
    {
        "id": "TC_AC_002", "category": "Happy Path — Add Client", "priority": "Critical",
        "title": "Add B2B Client with all mandatory fields — full happy path",
        "preconditions": "Add Client page is open. At least 1 Sales Partner exists.",
        "steps": "1. Select Client Source from dropdown\n2. Enter Client Name\n3. Select B2B\n4. Enter Email\n5. Enter 10-digit Mobile\n6. Type city & select from autocomplete\n7. Keep Active=Active\n8. Click 'Add Client'",
        "data": "Client Source: Direct | Client Name: Test Hotel Pvt Ltd | B2B: Yes | Email: test@hotel.com | Mobile: 9876543210 | City: Delhi | Active: Active",
        "expected": "1. Redirects to Client List page\n2. Flash message: 'Client record is updated successfully.'\n3. New record visible in list with Auto Generated ID\n4. All entered values match the record"
    },
    {
        "id": "TC_AC_003", "category": "Happy Path — Add Client", "priority": "Critical",
        "title": "Add B2C Client with optional Reference and Remarks",
        "preconditions": "Add Client page is open",
        "steps": "1. Select Client Source\n2. Enter Client Name\n3. Select B2C\n4. Enter Email\n5. Enter Mobile\n6. Select City\n7. Enter Reference\n8. Enter Remarks\n9. Click 'Add Client'",
        "data": "Client Name: Rahul Sharma | B2B: No (B2C) | Email: rahul@mail.com | Mobile: 9123456789 | City: Delhi | Reference: Google Ads | Remarks: VIP client",
        "expected": "Client saved with B2C type, Reference and Remarks stored correctly."
    },
    {
        "id": "TC_AC_004", "category": "Happy Path — Add Client", "priority": "High",
        "title": "Add Client with Active status = In-Active",
        "preconditions": "Add Client page is open",
        "steps": "1. Fill all mandatory fields\n2. Change Active dropdown to 'In-Active'\n3. Click 'Add Client'",
        "data": "Active: In-Active (false)",
        "expected": "Client saved. In list page, status shows 'In-Active' (red badge)."
    },
    # ── City Validation ──
    {
        "id": "TC_AC_005", "category": "City Validation", "priority": "Critical",
        "title": "Submit with city typed but NOT selected from autocomplete",
        "preconditions": "Add Client page loaded",
        "steps": "1. Fill all fields\n2. In City field, type 'Delhi' but do NOT click the autocomplete suggestion\n3. Click 'Add Client'",
        "data": "City text: Delhi (not selected from dropdown)",
        "expected": "Form shows error on City field (city.error). Form re-renders with all data retained. Client is NOT saved."
    },
    {
        "id": "TC_AC_006", "category": "City Validation", "priority": "High",
        "title": "Submit with invalid/non-existent city name",
        "preconditions": "Add Client page loaded",
        "steps": "1. Fill all fields\n2. Type 'XYZNonExistent' in city\n3. Click 'Add Client'",
        "data": "City: XYZNonExistent",
        "expected": "City field shows validation error. Form re-loads with Sales Partner dropdown repopulated. Client NOT saved."
    },
    # ── Security ──
    {
        "id": "TC_AC_007", "category": "Security & Authorization", "priority": "Critical",
        "title": "Access Add Client page without login",
        "preconditions": "User is NOT logged in. Browser session cleared.",
        "steps": "1. Open browser\n2. Directly navigate to /ashokacrm/view_add_client_form",
        "data": "URL: /ashokacrm/view_add_client_form",
        "expected": "Redirects to login page. Add Client page is NOT accessible."
    },
    # ── Edge Cases ──
    {
        "id": "TC_AC_008", "category": "Edge Cases & Data Integrity", "priority": "High",
        "title": "XSS payload in Client Name — verify it doesn't execute",
        "preconditions": "Add Client page loaded",
        "steps": "1. Enter '<script>alert(1)</script>' as Client Name\n2. Fill other fields\n3. Submit\n4. Go to Client List and view the record",
        "data": "Client Name: <script>alert(1)</script>",
        "expected": "No JavaScript popup. The text should be escaped/displayed as text in the listing. If script executes → SECURITY BUG."
    },
    {
        "id": "TC_AC_009", "category": "Edge Cases & Data Integrity", "priority": "High",
        "title": "SQL injection attempt in Client Name",
        "preconditions": "Add Client page loaded",
        "steps": "1. Enter \"' OR 1=1; DROP TABLE client; --\" as Client Name\n2. Fill other fields\n3. Submit and verify",
        "data": "Client Name: ' OR 1=1; DROP TABLE client; --",
        "expected": "Client saved as literal text. No DB error. Other records not affected. System functional."
    },
    {
        "id": "TC_AC_010", "category": "Edge Cases & Data Integrity", "priority": "Medium",
        "title": "Client Name with only whitespace",
        "preconditions": "Add Client page loaded",
        "steps": "1. Enter only spaces in Client Name\n2. Fill other fields\n3. Submit",
        "data": "Client Name: '     ' (5 spaces)",
        "expected": "Ideally: Validation error (name is required). Current behavior: saves as-is → marks as POTENTIAL BUG if whitespace-only is saved."
    },
    {
        "id": "TC_AC_011", "category": "Edge Cases & Data Integrity", "priority": "Medium",
        "title": "Duplicate client name + mobile — submit twice",
        "preconditions": "Add Client page loaded. A client 'Duplicate Hotel' already exists.",
        "steps": "1. Add client with Name='Duplicate Hotel', Mobile='9876543210'\n2. Submit\n3. Repeat with exact same data\n4. Check client list",
        "data": "Client Name: Duplicate Hotel (both times)",
        "expected": "Both records saved (no uniqueness check at controller level). IF duplicates should be blocked → BUG."
    },
    # ── Backend Logic ──
    {
        "id": "TC_AC_012", "category": "Backend & DB Verification", "priority": "High",
        "title": "Verify Client ID is auto-generated in DB",
        "preconditions": "Client successfully added",
        "steps": "1. Add a new client\n2. Go to Client List\n3. Verify the new record has an auto-generated ID",
        "data": "N/A",
        "expected": "Client ID is a unique auto-incremented integer. Not null."
    },
    {
        "id": "TC_AC_013", "category": "Backend & DB Verification", "priority": "High",
        "title": "Verify City and Sales Partner foreign keys are correctly stored",
        "preconditions": "Client added with City=Delhi, Source=Direct",
        "steps": "1. Add client selecting City=Delhi, Sales Partner=Direct\n2. Check DB record or client details page",
        "data": "City: Delhi | Sales Partner: Direct",
        "expected": "In DB: city_destination_id matches Delhi's ID. sales_partner_id matches Direct's ID. Data integrity maintained."
    },
    {
        "id": "TC_AC_014", "category": "Post-Submission", "priority": "High",
        "title": "Verify redirect to Client List and success flash message",
        "preconditions": "Client added successfully",
        "steps": "1. Add a valid client\n2. Observe the page after submission",
        "data": "N/A",
        "expected": "1. Page redirects to view_clients_list\n2. Green success message: 'Client record is updated successfully.'\n3. New client appears in the list"
    },
    {
        "id": "TC_AC_015", "category": "Post-Submission", "priority": "Medium",
        "title": "View Client List button navigates correctly",
        "preconditions": "On Add Client page",
        "steps": "1. Click 'View Client List' button (without submitting)",
        "data": "N/A",
        "expected": "Navigates to /ashokacrm/view_clients_list page showing all clients."
    },
]

# ═══════════════════════════════════════════════════════════════════════════
#  RATE TYPE MANAGEMENT — TEST CASES
#  Feature: Add Rate Type, Edit Rate Type, View Rate Type List
#  Module:  Sales Management
#  Endpoints:
#    GET  /view_add_rate_type_form  → Admin_Add_RateType.jsp
#    POST /create_edit_rate_type    → saves & redirects to list
#    GET  /view_rate_type_list      → viewRateTypeListing.jsp
#    POST /view_edit_rate_type_form → Admin_Edit_RateType.jsp
#    POST /view_rate_type_sessionwise → viewRateSessionMappingList.jsp
# ═══════════════════════════════════════════════════════════════════════════
RATE_TYPE_TESTS = [
    # ── Page Load — Add Rate Type ──
    {
        "id": "TC_RT_001", "category": "Add Rate Type — Page Load", "priority": "High",
        "title": "Verify Add Rate Type page loads successfully",
        "preconditions": "User is logged in as ADMIN",
        "steps": "1. Navigate to Sales Management menu\n2. Click 'Add Rate Type'",
        "data": "URL: /ashokacrm/view_add_rate_type_form",
        "expected": "Page loads with fields: Rate Type Id (Auto Generated in blue), Rate Type Name (text input, required), Description (textarea, max 500 chars). Buttons: 'Add Rate Type', 'View Rate Types List'. Background image loads correctly."
    },
    {
        "id": "TC_RT_002", "category": "Add Rate Type — Page Load", "priority": "Medium",
        "title": "Verify Rate Type Id shows 'Auto Generated' label",
        "preconditions": "Add Rate Type page is open",
        "steps": "1. Observe the 'Rate Type Id' field",
        "data": "N/A",
        "expected": "Rate Type Id displays 'Auto Generated' text in blue font. No input field shown for Rate Type Id."
    },
    # ── Happy Path — Add Rate Type ──
    {
        "id": "TC_RT_003", "category": "Add Rate Type — Happy Path", "priority": "Critical",
        "title": "Add Rate Type with valid name and description — full happy path",
        "preconditions": "Add Rate Type page is open",
        "steps": "1. Enter Rate Type Name\n2. Enter Description\n3. Click 'Add Rate Type'",
        "data": "Rate Type Name: Rack Rate | Description: Standard published rates for walk-in guests",
        "expected": "1. Redirects to Rate Type List page\n2. Flash message: 'Rate Type Record is updated Successfully..'\n3. New 'Rack Rate' record visible in table with auto-generated ID\n4. Status shows 'Active' (green badge)"
    },
    {
        "id": "TC_RT_004", "category": "Add Rate Type — Happy Path", "priority": "High",
        "title": "Add Rate Type with name only — no description",
        "preconditions": "Add Rate Type page is open",
        "steps": "1. Enter Rate Type Name\n2. Leave Description empty\n3. Click 'Add Rate Type'",
        "data": "Rate Type Name: Corporate Rate | Description: (empty)",
        "expected": "Rate Type saved successfully. Description stored as null/empty. Redirects to list. Record shows in table."
    },
    {
        "id": "TC_RT_005", "category": "Add Rate Type — Happy Path", "priority": "High",
        "title": "Add Rate Type with maximum length description (500 chars)",
        "preconditions": "Add Rate Type page is open",
        "steps": "1. Enter Rate Type Name\n2. Enter 500 characters in Description\n3. Click 'Add Rate Type'",
        "data": "Rate Type Name: Seasonal Rate | Description: 500 chars long text...",
        "expected": "Rate Type saved. Description stored fully (500 chars). No truncation."
    },
    # ── Validation — Add Rate Type ──
    {
        "id": "TC_RT_006", "category": "Add Rate Type — Validation", "priority": "Critical",
        "title": "Submit Add Rate Type with empty Rate Type Name",
        "preconditions": "Add Rate Type page is open",
        "steps": "1. Leave Rate Type Name empty\n2. Enter Description\n3. Click 'Add Rate Type'",
        "data": "Rate Type Name: (empty) | Description: Some description",
        "expected": "HTML5 required validation prevents submission. Browser shows 'Please fill out this field' tooltip on Rate Type Name."
    },
    {
        "id": "TC_RT_007", "category": "Add Rate Type — Validation", "priority": "High",
        "title": "Rate Type Name with only whitespace characters",
        "preconditions": "Add Rate Type page is open",
        "steps": "1. Enter only spaces in Rate Type Name\n2. Click 'Add Rate Type'",
        "data": "Rate Type Name: '    ' (spaces only)",
        "expected": "Ideally: Validation error. Current: may save whitespace-only name. If saved → mark as POTENTIAL BUG."
    },
    {
        "id": "TC_RT_008", "category": "Add Rate Type — Validation", "priority": "High",
        "title": "XSS payload in Rate Type Name",
        "preconditions": "Add Rate Type page is open",
        "steps": "1. Enter '<script>alert(\"XSS\")</script>' as Rate Type Name\n2. Submit\n3. View Rate Type List",
        "data": "Rate Type Name: <script>alert('XSS')</script>",
        "expected": "No script execution. Name displayed as escaped text in list. If alert pops up → SECURITY BUG."
    },
    {
        "id": "TC_RT_009", "category": "Add Rate Type — Validation", "priority": "Medium",
        "title": "SQL injection in Rate Type Name",
        "preconditions": "Add Rate Type page is open",
        "steps": "1. Enter SQL injection string as name\n2. Submit\n3. Verify system integrity",
        "data": "Rate Type Name: ' OR 1=1; DROP TABLE ratetype; --",
        "expected": "Saved as literal text. No DB error or data loss. JPA parameterized queries prevent injection."
    },
    {
        "id": "TC_RT_010", "category": "Add Rate Type — Validation", "priority": "Medium",
        "title": "Description exceeding 500 characters",
        "preconditions": "Add Rate Type page is open",
        "steps": "1. Enter Rate Type Name\n2. Try to type more than 500 chars in Description\n3. Check if textarea restricts input",
        "data": "Description: 501+ characters",
        "expected": "Textarea has maxlength=500. Characters beyond 500 are not accepted. Verified in browser."
    },
    # ── Edit Rate Type ──
    {
        "id": "TC_RT_011", "category": "Edit Rate Type", "priority": "Critical",
        "title": "Edit Rate Type — update name and description",
        "preconditions": "At least 1 Rate Type exists in list",
        "steps": "1. Go to Rate Type List\n2. Click 'Edit' button on a record\n3. Edit page loads with current values\n4. Change Rate Type Name and Description\n5. Click 'Edit Rate Type'",
        "data": "Original Name: Rack Rate → New Name: Premium Rack Rate | Description updated",
        "expected": "1. Edit page shows current values pre-filled\n2. Rate Type Id displayed (not editable)\n3. Active dropdown visible (Active/In-Active)\n4. After save: redirect to list with message 'Rate Type Record is updated Successfully..'\n5. Updated name visible in list"
    },
    {
        "id": "TC_RT_012", "category": "Edit Rate Type", "priority": "High",
        "title": "Edit Rate Type — change Active to In-Active",
        "preconditions": "An active Rate Type exists",
        "steps": "1. Go to Rate Type List\n2. Click 'Edit' on an active Rate Type\n3. Change Active dropdown to 'In-Active'\n4. Click 'Edit Rate Type'",
        "data": "Active: In-Active",
        "expected": "Rate Type status changes to 'In-Active' (red badge in list). This Rate Type should NOT appear in Sales Partner dropdown."
    },
    {
        "id": "TC_RT_013", "category": "Edit Rate Type", "priority": "High",
        "title": "Edit Rate Type — change In-Active back to Active",
        "preconditions": "An in-active Rate Type exists",
        "steps": "1. Go to Rate Type List\n2. Click 'Edit' on an in-active Rate Type\n3. Change Active dropdown to 'Active'\n4. Click 'Edit Rate Type'",
        "data": "Active: Active",
        "expected": "Rate Type status changes to 'Active' (green badge). Now appears in Sales Partner Rate Type dropdown again."
    },
    {
        "id": "TC_RT_014", "category": "Edit Rate Type", "priority": "Medium",
        "title": "Edit Rate Type — submit without changing anything",
        "preconditions": "Edit Rate Type page loaded with existing data",
        "steps": "1. Click 'Edit' on a Rate Type\n2. Do NOT change anything\n3. Click 'Edit Rate Type'",
        "data": "No changes",
        "expected": "Record saved without error. Data unchanged. Success message shown. No new duplicate created."
    },
    # ── View Rate Type List ──
    {
        "id": "TC_RT_015", "category": "View Rate Type List", "priority": "High",
        "title": "Verify Rate Type List page loads and displays all records",
        "preconditions": "User is logged in. At least 2 Rate Types exist.",
        "steps": "1. Navigate to Rate Type List page",
        "data": "URL: /ashokacrm/view_rate_type_list",
        "expected": "Table displays columns: Rate Type ID, Rate Type Name, Description, Status, Action. All Rate Types shown. Active=green badge, In-Active=red badge."
    },
    {
        "id": "TC_RT_016", "category": "View Rate Type List", "priority": "High",
        "title": "Edit button navigates to Edit Rate Type form",
        "preconditions": "On Rate Type List page with records",
        "steps": "1. Click 'Edit' button on any record",
        "data": "N/A",
        "expected": "Navigates to Edit Rate Type form. Form is pre-filled with selected Rate Type's data. Rate Type Id shown in blue."
    },
    {
        "id": "TC_RT_017", "category": "View Rate Type List", "priority": "Medium",
        "title": "'Check Applicable Dates' button works",
        "preconditions": "On Rate Type List page. Rate Type has session-rate mappings.",
        "steps": "1. Click '📅 Check Applicable Dates' button on a record",
        "data": "N/A",
        "expected": "Navigates to Rate Session Mapping List. Shows date ranges linked to this Rate Type with formatted start/end dates."
    },
    {
        "id": "TC_RT_018", "category": "View Rate Type List", "priority": "Medium",
        "title": "Empty Rate Type List — no records in DB",
        "preconditions": "No Rate Types exist in system (or all deleted)",
        "steps": "1. Navigate to Rate Type List",
        "data": "N/A",
        "expected": "Page loads with empty table. Table headers visible but no data rows. No error."
    },
    # ── Security ──
    {
        "id": "TC_RT_019", "category": "Security", "priority": "Critical",
        "title": "Access Add Rate Type page without login",
        "preconditions": "User is NOT logged in",
        "steps": "1. Open browser\n2. Navigate directly to /ashokacrm/view_add_rate_type_form",
        "data": "URL: /ashokacrm/view_add_rate_type_form",
        "expected": "Redirects to login page. Page is NOT accessible."
    },
    {
        "id": "TC_RT_020", "category": "Security", "priority": "High",
        "title": "POST to create_edit_rate_type without login",
        "preconditions": "User is NOT logged in",
        "steps": "1. Use Postman/cURL to POST to /ashokacrm/create_edit_rate_type with rate type data",
        "data": "POST body: rateTypeName=Test&description=Test",
        "expected": "Returns 302 redirect to login page. Rate Type NOT saved."
    },
    # ── Integration ──
    {
        "id": "TC_RT_021", "category": "Integration with Sales Partner", "priority": "High",
        "title": "New active Rate Type appears in Sales Partner form dropdown",
        "preconditions": "A new active Rate Type 'Weekend Rate' was just added",
        "steps": "1. Add Rate Type 'Weekend Rate' (Active)\n2. Navigate to Add Sales Partner form\n3. Check Rate Type dropdown",
        "data": "Rate Type: Weekend Rate",
        "expected": "'Weekend Rate' appears in the Rate Type dropdown on the Sales Partner form."
    },
    {
        "id": "TC_RT_022", "category": "Integration with Sales Partner", "priority": "High",
        "title": "In-active Rate Type does NOT appear in Sales Partner dropdown",
        "preconditions": "A Rate Type is set to In-Active",
        "steps": "1. Edit a Rate Type → set Active to In-Active\n2. Navigate to Add Sales Partner form\n3. Check Rate Type dropdown",
        "data": "Rate Type status: In-Active",
        "expected": "The in-active Rate Type does NOT appear in the dropdown. Only active Rate Types shown."
    },
]


# ═══════════════════════════════════════════════════════════════════════════
#  GENERATE ALL FILES
# ═══════════════════════════════════════════════════════════════════════════
OUTPUT_DIR = "/Users/sarthakjain/vistaluxhms_repository/base_vistaluxhms_repository/test_cases"

# 1. Add Client — Shortlisted
ac_excel = os.path.join(OUTPUT_DIR, "Add_Client_Test_Cases.xlsx")
ac_csv   = os.path.join(OUTPUT_DIR, "Add_Client_Test_Cases.csv")
ac_count = write_excel(ADD_CLIENT_TESTS, ac_excel, "Add Client Tests")
write_csv(ADD_CLIENT_TESTS, ac_csv)
print(f"✅  Add Client:   {ac_count} essential test cases")
print(f"    Excel: {ac_excel}")
print(f"    CSV:   {ac_csv}")

# 2. Rate Type Management
rt_excel = os.path.join(OUTPUT_DIR, "Rate_Type_Test_Cases.xlsx")
rt_csv   = os.path.join(OUTPUT_DIR, "Rate_Type_Test_Cases.csv")
rt_count = write_excel(RATE_TYPE_TESTS, rt_excel, "Rate Type Tests")
write_csv(RATE_TYPE_TESTS, rt_csv)
print(f"\n✅  Rate Type:     {rt_count} test cases")
print(f"    Excel: {rt_excel}")
print(f"    CSV:   {rt_csv}")

print(f"\n📊  Grand Total:   {ac_count + rt_count} test cases across 2 feature files")
print(f"📂  Output folder: {OUTPUT_DIR}")
